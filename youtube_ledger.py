"""本地台账、公开页面字段补齐和参考 Excel 导出。"""
from __future__ import annotations

import copy
import json
import re
import sqlite3
import urllib.request
import uuid
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path

HEADERS = ["视频内容", "视频链接", "发布时间", "观看量", "点赞量", "评论量", "频道名称", "网红频道链接", "粉丝数", "地区", "类型"]
FIELDS = ["title", "video_url", "published_at", "views", "likes", "comments", "channel_name", "channel_url", "subscribers", "region", "video_type"]
LOCAL_TZ = timezone(timedelta(hours=8))


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def local_stamp():
    return datetime.now(LOCAL_TZ).strftime("%Y%m%d_%H%M%S")


def walk_named(value, key):
    if isinstance(value, dict):
        if key in value:
            yield value[key]
        for child in value.values():
            yield from walk_named(child, key)
    elif isinstance(value, list):
        for child in value:
            yield from walk_named(child, key)


def text_of(value):
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        if "simpleText" in value:
            return value["simpleText"]
        if "content" in value:
            return value["content"]
        return "".join(item.get("text", "") for item in value.get("runs", []))
    return ""


def parse_count(value):
    text = text_of(value).replace(",", "").replace("\u00a0", " ")
    match = re.search(r"(?<!\w)(\d+(?:\.\d+)?)\s*([KMB])?", text, re.I)
    if not match:
        return None
    multipliers = {"": 1, "K": 1000, "M": 1000000, "B": 1000000000}
    return int(float(match[1]) * multipliers.get((match[2] or "").upper(), 1))


def json_after(html, pattern):
    match = re.search(pattern, html)
    if not match:
        return {}
    try:
        return json.JSONDecoder().raw_decode(html[match.end():].lstrip())[0]
    except ValueError:
        return {}


def page_data(url):
    request = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/130.0.0.0 Safari/537.36",
        "Accept-Language": "en-US,en;q=0.9",
    })
    with urllib.request.urlopen(request, timeout=20) as response:
        html = response.read().decode("utf-8", "replace")
    if "unusual traffic" in html.lower():
        raise RuntimeError("YouTube 网页要求验证访问状态")
    initial = json_after(html, r"(?:var\s+)?ytInitialData\s*=\s*")
    config = {}
    for match in re.finditer(r"ytcfg\.set\(", html):
        try:
            value = json.JSONDecoder().raw_decode(html[match.end():].lstrip())[0]
            if isinstance(value, dict):
                config.update(value)
        except ValueError:
            continue
    player = json_after(html, r"(?:var\s+)?ytInitialPlayerResponse\s*=\s*")
    return initial, config, player


def continuation_data(endpoint, token, config):
    context = config.get("INNERTUBE_CONTEXT")
    if not context:
        return {}
    payload = json.dumps({"context": context, "continuation": token}).encode("utf-8")
    request = urllib.request.Request("https://www.youtube.com/youtubei/v1/" + endpoint,
        data=payload, headers={"Content-Type": "application/json", "User-Agent": "Mozilla/5.0", "Origin": "https://www.youtube.com"})
    with urllib.request.urlopen(request, timeout=20) as response:
        return json.load(response)


def country_code(name):
    aliases = {"South Korea": "KR", "North Korea": "KP", "Russia": "RU", "Taiwan": "TW", "Vietnam": "VN", "Bolivia": "BO", "Venezuela": "VE", "Tanzania": "TZ", "Moldova": "MD", "Iran": "IR", "Syria": "SY", "Laos": "LA", "Kosovo": "XK"}
    if name in aliases:
        return aliases[name]
    if re.fullmatch(r"[A-Z]{2}", name or ""):
        return name
    try:
        import pycountry
        return pycountry.countries.lookup(name).alpha_2
    except (LookupError, ImportError):
        return name or "未公开"


def channel_details(channel_id):
    url = f"https://www.youtube.com/channel/{channel_id}/about?hl=en"
    data, config, _ = page_data(url)
    models = list(walk_named(data, "aboutChannelViewModel"))
    old_models = list(walk_named(data, "channelAboutFullMetadataRenderer"))
    if not models and not old_models:
        for panel in walk_named(data, "showEngagementPanelEndpoint"):
            commands = list(walk_named(panel, "continuationCommand"))
            if commands:
                more = continuation_data("browse", commands[0]["token"], config)
                models = list(walk_named(more, "aboutChannelViewModel"))
                if models:
                    break
    for model in models + old_models:
        if model.get("channelId") and model["channelId"] != channel_id:
            continue
        country = text_of(model.get("country"))
        return {
            "region": country_code(country) if country else "未公开",
            "country_name": country, "subscribers": parse_count(model.get("subscriberCountText")),
            "channel_description": text_of(model.get("description")),
            "region_source": url, "collected_at": now_iso(),
        }
    raise ValueError("频道关于页结构未识别")


def page_metrics(video_id):
    data, config, player = page_data(f"https://www.youtube.com/watch?v={video_id}&hl=en")
    details = player.get("videoDetails", {})
    micro = player.get("microformat", {}).get("playerMicroformatRenderer", {})
    result = {"views": parse_count(details.get("viewCount", "")),
              "published_at": (micro.get("publishDate") or micro.get("uploadDate") or "")[:10],
              "title": details.get("title"), "channel_name": details.get("author"), "channel_id": details.get("channelId")}
    for view in walk_named(data, "viewCount"):
        count = parse_count(view)
        if count is not None and result["views"] is None:
            result["views"] = count
    for label in walk_named(data, "label"):
        if isinstance(label, str) and re.search(r"\d[\d,.KMB ]*\s+likes?\b", label, re.I):
            result["likes"] = parse_count(label)
            break
    for key in ("commentCount", "commentsHeaderRenderer", "commentsEntryPointHeaderRenderer"):
        for item in walk_named(data, key):
            value = item.get("countText", item.get("commentCount")) if isinstance(item, dict) else item
            count = parse_count(value)
            if count is not None:
                result["comments"] = count
                break
        if result.get("comments") is not None:
            break
    if result.get("comments") is None:
        for section in walk_named(data, "itemSectionRenderer"):
            if section.get("sectionIdentifier") != "comment-item-section":
                continue
            commands = list(walk_named(section, "continuationCommand"))
            if commands:
                more = continuation_data("next", commands[0]["token"], config)
                for header in walk_named(more, "commentsHeaderRenderer"):
                    count = parse_count(header.get("countText"))
                    if count is not None:
                        result["comments"] = count
                        break
                break
    # 只有页面明确写出关闭或暂无评论时才使用 0。
    if result.get("comments") is None:
        for message in walk_named(data, "messageRenderer"):
            text = text_of(message.get("text")).lower()
            if "comments are turned off" in text or "no comments yet" in text:
                result["comments"] = 0
                break
    return result


class Ledger:
    def __init__(self, path, legacy_path=None):
        self.path = Path(path)
        with self.connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS settings (name TEXT PRIMARY KEY, value TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS videos (video_id TEXT PRIMARY KEY, data TEXT NOT NULL, deleted INTEGER NOT NULL DEFAULT 0, first_seen TEXT, updated_at TEXT);
                CREATE TABLE IF NOT EXISTS detail_cache (video_id TEXT PRIMARY KEY, data TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS channels (channel_id TEXT PRIMARY KEY, data TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS runs (run_id TEXT PRIMARY KEY, started_at TEXT, finished_at TEXT, label TEXT, config TEXT);
                CREATE TABLE IF NOT EXISTS run_items (run_id TEXT, video_id TEXT, snapshot TEXT, PRIMARY KEY(run_id,video_id));
                CREATE INDEX IF NOT EXISTS run_items_video ON run_items(video_id);
            """)
            if "snapshot" not in [row[1] for row in db.execute("PRAGMA table_info(run_items)")]:
                db.execute("ALTER TABLE run_items ADD COLUMN snapshot TEXT")
        if legacy_path and Path(legacy_path).exists():
            legacy = Path(legacy_path)
            signature = [legacy.stat().st_mtime_ns, legacy.stat().st_size]
            if self.setting("legacy_signature") != signature:
                self.import_legacy(legacy)
                self.set_setting("legacy_signature", signature)

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.path, timeout=20)
        db.execute("PRAGMA journal_mode=WAL")
        try:
            yield db
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    def setting(self, key, default=None):
        with self.connect() as db:
            row = db.execute("SELECT value FROM settings WHERE name=?", (key,)).fetchone()
        return json.loads(row[0]) if row else default

    def set_setting(self, key, value):
        with self.connect() as db:
            db.execute("INSERT OR REPLACE INTO settings VALUES (?,?)", (key, json.dumps(value, ensure_ascii=False)))

    def create_run(self, config, label="采集"):
        run_id = local_stamp() + "_" + uuid.uuid4().hex[:6]
        with self.connect() as db:
            db.execute("INSERT INTO runs VALUES (?,?,NULL,?,?)", (run_id, now_iso(), label, json.dumps(config, ensure_ascii=False)))
        self.set_setting("last_run", run_id)
        return run_id

    def finish_run(self, run_id):
        with self.connect() as db:
            db.execute("UPDATE runs SET finished_at=? WHERE run_id=?", (now_iso(), run_id))

    def runs(self):
        with self.connect() as db:
            rows = db.execute("SELECT run_id,started_at,label FROM runs ORDER BY started_at DESC").fetchall()
        return [{"run_id": r[0], "started_at": r[1], "label": r[2]} for r in rows]

    def get_video(self, video_id):
        with self.connect() as db:
            row = db.execute("SELECT data,deleted FROM videos WHERE video_id=?", (video_id,)).fetchone()
        if row:
            data = json.loads(row[0])
            data["_deleted"] = bool(row[1])
            return data
        with self.connect() as db:
            row = db.execute("SELECT data FROM detail_cache WHERE video_id=?", (video_id,)).fetchone()
        return json.loads(row[0]) if row else None

    def store_cache(self, data):
        with self.connect() as db:
            db.execute("INSERT OR REPLACE INTO detail_cache VALUES (?,?)", (data["video_id"], json.dumps(data, ensure_ascii=False)))

    def channel(self, channel_id):
        with self.connect() as db:
            row = db.execute("SELECT data FROM channels WHERE channel_id=?", (channel_id,)).fetchone()
        return json.loads(row[0]) if row else None

    def store_channel(self, channel_id, data):
        with self.connect() as db:
            db.execute("INSERT OR REPLACE INTO channels VALUES (?,?)", (channel_id, json.dumps(data, ensure_ascii=False)))

    def add_video(self, row, run_id=None):
        row = copy.deepcopy(row)
        video_id = row["video_id"]
        with self.connect() as db:
            found = db.execute("SELECT data,deleted,first_seen FROM videos WHERE video_id=?", (video_id,)).fetchone()
            if found and found[1]:
                return False
            old = json.loads(found[0]) if found else {}
            manual = old.get("_manual", {})
            brands = list(dict.fromkeys(old.get("brands", [old.get("brand")]) + row.get("brands", []) + [row.get("brand")]))
            for key, value in old.items():
                missing = (None, "", "未公开", "获取失败")
                if manual.get(key) or row.get(key) in missing and value not in missing or key not in row:
                    row[key] = value
            row["brands"] = [name for name in brands if name]
            row["_manual"] = manual
            row["key"] = video_id
            row.pop("status", None)
            row.pop("_deleted", None)
            first_seen = found[2] if found else row.get("collected_at", now_iso())
            db.execute("INSERT OR REPLACE INTO videos VALUES (?,?,0,?,?)", (video_id, json.dumps(row, ensure_ascii=False), first_seen, now_iso()))
            if run_id:
                db.execute("INSERT OR REPLACE INTO run_items (run_id,video_id,snapshot) VALUES (?,?,?)", (run_id, video_id, json.dumps(row, ensure_ascii=False)))
        return True

    def rows(self, run_id=None, deleted=False):
        with self.connect() as db:
            if run_id:
                records = db.execute("SELECT COALESCE(r.snapshot,v.data) FROM videos v JOIN run_items r USING(video_id) WHERE r.run_id=? AND v.deleted=?", (run_id, int(deleted))).fetchall()
            else:
                records = db.execute("SELECT data FROM videos WHERE deleted=?", (int(deleted),)).fetchall()
        return sorted([json.loads(r[0]) for r in records], key=lambda row: row.get("published_at", ""), reverse=True)

    def set_deleted(self, ids, deleted=True):
        with self.connect() as db:
            db.executemany("UPDATE videos SET deleted=? WHERE video_id=?", [(int(deleted), video_id) for video_id in ids])

    def edit(self, video_id, region, video_type):
        row = self.get_video(video_id)
        if not row:
            return
        row["region"] = region.strip() or "未公开"
        row["video_type"] = video_type.strip() or "Other"
        row.setdefault("_manual", {}).update({"region": True, "video_type": True})
        row.pop("_deleted", None)
        with self.connect() as db:
            db.execute("UPDATE videos SET data=? WHERE video_id=?", (json.dumps(row, ensure_ascii=False), video_id))
            for run_id, snapshot in db.execute("SELECT run_id,snapshot FROM run_items WHERE video_id=?", (video_id,)).fetchall():
                item = json.loads(snapshot) if snapshot else dict(row)
                item["region"], item["video_type"] = row["region"], row["video_type"]
                item["_manual"] = dict(row["_manual"])
                db.execute("UPDATE run_items SET snapshot=? WHERE run_id=? AND video_id=?", (json.dumps(item, ensure_ascii=False), run_id, video_id))

    def import_legacy(self, path):
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        brands = []
        for line in data.get("brands_text", "").splitlines():
            parts = [part.strip() for part in line.split("|", 2)]
            if len(parts) == 3:
                brands.append({"name": parts[0], "official": parts[1], "terms": parts[2]})
        if brands and not self.setting("brands"):
            self.set_setting("brands", brands)
        records = data.get("candidates", {})
        with self.connect() as db:
            existing = {item[0] for item in db.execute("SELECT video_id FROM videos")}
        new_records = [row for row in records.values() if row.get("video_id") and row["video_id"] not in existing]
        if new_records:
            run_id = self.create_run({}, "旧版历史导入")
            for row in new_records:
                self.add_video(row, run_id)
                if row.get("status") == "已排除":
                    self.set_deleted([row["video_id"]])
            self.finish_run(run_id)
        for row in data.get("web_cache", {}).values():
            if row.get("video_id") and not self.get_video(row["video_id"]):
                self.store_cache(row)
        self.set_setting("legacy_imported", True)


def table_values(row):
    values = []
    for field in FIELDS:
        value = row.get(field)
        if field == "published_at":
            value = (value or "")[:10]
        if value is None or value == "":
            value = "未公开" if field in ("views", "likes", "comments", "subscribers", "region") else ""
        values.append(value)
    return values


def export_reference(rows, output_path, template_path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    # 开源分发不携带业务工作簿；本地自定义模板仍可覆盖默认样式。
    if Path(template_path).is_file():
        template = load_workbook(template_path)
    else:
        template = Workbook()
        default = template.active
        default.title = "视频数据"
        default.append(HEADERS)
        for cell in default[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="27445C")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for cell in default[2]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        default.cell(2, 3).number_format = "yyyy/m/d"
    source = template.active
    if [source.cell(1, i).value for i in range(1, 12)] != HEADERS:
        template.close()
        raise ValueError("参考 Excel 的 11 列表头与要求不一致")
    book = Workbook()
    sheet = book.active
    sheet.title = source.title
    sheet.sheet_format = copy.copy(source.sheet_format)
    for i, header in enumerate(HEADERS, 1):
        sheet.cell(1, i, header)
    def apply_style(target, ref):
        target.font = copy.copy(ref.font)
        target.fill = copy.copy(ref.fill)
        target.border = copy.copy(ref.border)
        target.alignment = copy.copy(ref.alignment)
        target.protection = copy.copy(ref.protection)
        target.number_format = ref.number_format
    for i in range(1, 12):
        apply_style(sheet.cell(1, i), source.cell(1, i))
    seen = set()
    for row in rows:
        if row["video_id"] in seen:
            continue
        seen.add(row["video_id"])
        values = table_values(row)
        if row.get("published_at"):
            try:
                values[2] = datetime.strptime(row["published_at"][:10], "%Y-%m-%d")
            except ValueError:
                pass
        sheet.append(values)
        # 将外部标题等字段明确写为文本，避免以 = 开头的内容成为公式。
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for i in range(1, 12):
            apply_style(sheet.cell(sheet.max_row, i), source.cell(2, i))
        for i in (2, 8):
            if values[i-1]:
                sheet.cell(sheet.max_row, i).hyperlink = values[i-1]
    for i, width in enumerate((64, 47, 15, 14, 14, 14, 25, 52, 14, 12, 17), 1):
        col = get_column_letter(i)
        sheet.column_dimensions[col].width = source.column_dimensions[col].width if col in source.column_dimensions else width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(output_path)
    template.close()
    book.close()
