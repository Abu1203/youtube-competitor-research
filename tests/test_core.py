import copy
from datetime import datetime
import importlib.util
import io
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
import urllib.error

from openpyxl import Workbook, load_workbook
from youtube_ledger import HEADERS, Ledger, export_reference

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("research_app", ROOT / "竞品红人视频挖掘.py")
app = importlib.util.module_from_spec(spec)
spec.loader.exec_module(app)


def sample():
    return {
        "video_id": "demo-video-1", "brand": "Example Brand", "title": "Example review",
        "published_at": "2026-01-02", "views": 0, "likes": None, "comments": 3,
        "video_url": "https://example.com/video/1", "channel_url": "https://example.com/channel/1",
        "channel_name": "Example Channel", "subscribers": 100, "region": "US", "video_type": "Review",
    }


class LedgerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.ledger = Ledger(self.root / "test.sqlite3")

    def test_deduplication_keeps_batch_snapshots(self):
        first = self.ledger.create_run({})
        second = self.ledger.create_run({})
        row = sample()
        self.ledger.add_video(row, first)
        row.update(views=22, brand="Another Example")
        self.ledger.add_video(row, second)
        self.assertEqual(len(self.ledger.rows()), 1)
        self.assertEqual(self.ledger.rows(first)[0]["views"], 0)
        self.assertEqual(self.ledger.rows(second)[0]["views"], 22)
        self.assertEqual(set(self.ledger.rows()[0]["brands"]), {"Example Brand", "Another Example"})

    def test_delete_blocks_recollection_until_restored(self):
        self.ledger.add_video(sample())
        self.ledger.set_deleted([sample()["video_id"]])
        self.assertEqual(self.ledger.rows(), [])
        self.assertFalse(self.ledger.add_video(sample()))
        self.assertEqual(len(self.ledger.rows(deleted=True)), 1)
        self.ledger.set_deleted([sample()["video_id"]], False)
        self.assertEqual(len(self.ledger.rows()), 1)

    def test_manual_fields_survive_refresh(self):
        run = self.ledger.create_run({})
        self.ledger.add_video(sample(), run)
        self.ledger.edit(sample()["video_id"], "FR", "DIY")
        row = sample()
        row["views"] = 19
        self.ledger.add_video(row, run)
        saved = self.ledger.rows(run)[0]
        self.assertEqual((saved["region"], saved["video_type"], saved["views"]), ("FR", "DIY", 19))

    def test_fresh_ledger_contains_no_brand_configuration(self):
        self.assertEqual(self.ledger.rows(), [])
        self.assertIsNone(self.ledger.setting("brands"))

    def test_export_without_private_template(self):
        output = self.root / "export.xlsx"
        row = sample()
        row["title"] = "=1+1"
        export_reference([row, copy.deepcopy(row)], output, self.root / "missing.xlsx")
        book = load_workbook(output)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
        self.assertEqual((sheet.max_row, sheet.max_column), (2, 11))
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertEqual(sheet["A2"].value, "=1+1")
        self.assertIsInstance(sheet["C2"].value, datetime)
        self.assertEqual(sheet["C2"].number_format, "yyyy/m/d")
        self.assertEqual(sheet["D2"].value, 0)
        self.assertEqual(sheet["E2"].value, "未公开")
        self.assertEqual(sheet["B2"].hyperlink.target, row["video_url"])
        self.assertEqual(sheet.freeze_panes, "A2")

    def test_custom_template_validation(self):
        template = self.root / "local.xlsx"
        book = Workbook()
        book.active.append(["invalid header"])
        book.save(template)
        book.close()
        with self.assertRaisesRegex(ValueError, "11 列"):
            export_reference([sample()], self.root / "output.xlsx", template)


class ApiTests(unittest.TestCase):
    def test_http_error_redacts_input_key(self):
        key = "synthetic-test-value"
        error = urllib.error.HTTPError("https://example.com", 400, "Bad Request", {}, io.BytesIO(
            ('{"error":{"message":"invalid ' + key + '"}}').encode()))
        with patch.object(app.urllib.request, "urlopen", side_effect=error):
            with self.assertRaises(app.ApiError) as caught:
                app.YouTubeApi(key, threading.Event()).get("search")
        self.assertNotIn(key, str(caught.exception))
        self.assertIn("[REDACTED]", str(caught.exception))

    def test_network_error_redacts_input_key(self):
        key = "synthetic-test-value"
        with patch.object(app.urllib.request, "urlopen", side_effect=urllib.error.URLError("key=" + key)):
            with patch.object(app.time, "sleep"):
                with self.assertRaises(app.ApiError) as caught:
                    app.YouTubeApi(key, threading.Event()).get("search")
        self.assertNotIn(key, str(caught.exception))


if __name__ == "__main__":
    unittest.main()
